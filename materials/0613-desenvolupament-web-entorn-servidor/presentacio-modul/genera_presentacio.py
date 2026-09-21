"""Genera el PPTX editable de presentació del mòdul.

Dependències: python-pptx i openpyxl. Execució des de qualsevol directori.
Les dades curriculars es llegeixen del full original; no es modifica.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
SOURCE = ROOT / 'raw/sources/IFC33C - Desenvolupament web en entorn servidor.xlsx'
LOGO = ROOT / 'skills/presentacions-pau-casesnoves/assets/logo-pau-casesnoves.png'
W = openpyxl.load_workbook(SOURCE, data_only=True)
R = {}
for row in range(4, 13):
    sh = W['2-ANCORATGE CURRICULAR']
    R[sh[f'A{row}'].value] = {
        'title': ' '.join(sh[f'B{row}'].value.split()),
        'weight': int(sh[f'C{row}'].value),
        'unit': int(sh[f'F{row}'].value),
    }
H = [int(W['6-Distribució temporal'][f'L{i}'].value) for i in range(5, 9)]
assert sum(H) == 230 and sum(r['weight'] for r in R.values()) == 100

GREEN, LIME, INK, MUTED, PALE, WHITE = '175F16', '32CD33', '202124', '595959', 'F3F5F3', 'FFFFFF'
FONT = 'Liberation Sans'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.author = 'David Pons'
prs.core_properties.title = 'Desenvolupament web en entorn servidor — Presentació del mòdul'
prs.core_properties.subject = 'DAW · 0613 · Curs 2026–2027'
prs.core_properties.language = 'ca-ES'
transcript = []


def box(s, x, y, w, h, color=PALE):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh


def text(s, txt, x, y, w, h, size=24, color=INK, bold=False, align=None):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.01)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(txt.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(6)
        if align is not None: p.alignment = align
    return sh


def slide(title, section, note=''):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.334, .09, GREEN)
    text(s, section.upper(), .58, .30, 12, .28, 11, GREEN, True)
    text(s, title, .58, .80, 12.15, 1.5 if len(prs.slides) == 1 else .95,
         40 if len(prs.slides) == 1 else 32, GREEN, True)
    box(s, .58, 7.03, 12.17, .012, 'DDE5DD')
    text(s, 'David Pons · CIFP Pau Casesnoves', .58, 7.12, 10, .24, 11, MUTED)
    text(s, f'{len(prs.slides):02}', 11.9, 7.10, .8, .27, 11, MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = note
    transcript.append((s, title, note))
    return s


def card(s, x, y, w, h, title, body, num=None):
    box(s, x, y, w, h)
    box(s, x, y, .045, h, GREEN)
    offset = .24
    if num:
        text(s, num, x + .24, y + .15, w - .48, .5, 28, GREEN, True)
        offset = .77
    text(s, title, x + .24, y + offset, w - .48, .73, 23, GREEN, True)
    text(s, body, x + .24, y + offset + .83, w - .48, h - offset - .9, 22)


def banner(s, txt, y=6.15):
    box(s, .58, y, 12.17, .64, GREEN)
    text(s, txt, .82, y + .10, 11.7, .43, 22, WHITE, True)


def ra(s, key, y, h=1.2, size=22):
    r = R[key]
    box(s, .58, y, 12.17, h)
    text(s, key, .80, y + .16, 1, .43, 22, GREEN, True)
    text(s, r['title'], 1.95, y + .12, 8.95, h - .16, size)
    text(s, f"{r['weight']}%", 11.40, y + .15, 1.05, .5, 26, GREEN, True, PP_ALIGN.RIGHT)


def unit_note(n):
    return (f'Font: {SOURCE.name}, full 6-Distribució temporal, fila {n+4}; '
            'full 2-ANCORATGE CURRICULAR, A4:F12. Títols dels RA extrets del full amb espais normalitzats. '
            'Els percentatges indiquen el pes de cada RA en la nota global del mòdul. '
            'Els resums d’unitat incorporen l’enfocament docent de wiki/orientacio-docent.md.')


s = slide('Desenvolupament web\nen entorn servidor', '0613 · DAW · 2026–2027',
          'Presentació del mòdul. Font horària: 1-INTRODUCCIÓ!C7,F7. Professor: David Pons.')
text(s, 'La lògica que fa funcionar\nuna aplicació web', .6, 2.8, 8.3, 1.5, 34, INK)
text(s, 'David Pons', .6, 4.8, 8, .65, 28, GREEN, True)
text(s, 'CIFP Pau Casesnoves', .6, 5.52, 8, .5, 24)
s.shapes.add_picture(str(LOGO), Inches(9.4), Inches(2.85), width=Inches(2.7))
banner(s, '230 hores · 7 hores setmanals · Java i Spring')

s = slide('Què passa quan fas clic a «Reservar»?', '01 · Una idea general',
          'Exemple docent il·lustratiu: reserva d’un llibre. No descriu un projecte imposat al curs.')
text(s, 'El navegador envia una petició. El servidor decideix què pot fer i retorna una resposta.', .6, 1.83, 12, .95, 25)
for x, title, body in [(0.6,'Navegador','L’usuari demana\nreservar un llibre.'),(4.85,'Servidor','Comprova disponibilitat\ni aplica les regles.'),(9.1,'Dades','Consulta i desa\nla reserva.')]:
    card(s, x, 3.05, 3.63, 2.35, title, body)
for x in [4.3, 8.55]: text(s, '→', x, 3.9, .48, .6, 28, GREEN, True)
banner(s, 'En aquest mòdul construirem la part del servidor.')

s = slide('Aprendre a construir aplicacions ben organitzades', '01 · Una idea general',
          'Font: wiki/orientacio-docent.md i CONTEXT_ASSIGNATURA.md del ZIP docent.')
card(s, .6, 2.0, 3.86, 3.55, 'Responsabilitats clares', 'Cada classe i cada capa tenen una funció dins l’aplicació.', '01')
card(s, 4.74, 2.0, 3.86, 3.55, 'Dades i serveis', 'Consultam dades, aplicam regles de negoci i exposam funcionalitats.', '02')
card(s, 8.88, 2.0, 3.86, 3.55, 'Qualitat del codi', 'Cercam aplicacions comprensibles, mantenibles i comprovables.', '03')
banner(s, 'Java i Spring seran les eines habituals del curs.')

context_source = ('Relacions pedagògiques sintetitzades per a aquesta presentació. Noms dels mòduls contrastats amb '
    'RD 686/2010, https://www.boe.es/buscar/doc.php?id=BOE-A-2010-9269, i RD 405/2023, '
    'https://www.boe.es/eli/es/rd/2023/05/29/405. No s’atribueixen hores ni ponderacions als altres mòduls.')
s = slide('Els fonaments que connectarem', '02 · El mòdul dins DAW', context_source)
for x,y,t,b in [(.6,1.95,'Programació','Classes, objectes, interfícies i estructures de control.'),(6.83,1.95,'Bases de dades','Model de dades, consultes i persistència.'),(.6,4.12,'Llenguatges de marques','HTML i formats d’intercanvi d’informació.'),(6.83,4.12,'Entorns de desenvolupament','Eines per escriure, depurar i provar el codi.')]:
    card(s,x,y,5.90,1.98,t,b)

s = slide('D’una interfície a una aplicació completa', '02 · El mòdul dins DAW', context_source)
card(s,.6,1.95,3.86,3.1,'Client i interfícies','Interacció al navegador, formularis i presentació de la informació.')
card(s,4.74,1.95,3.86,3.1,'Entorn servidor','Peticions HTTP, regles de negoci, dades i serveis web.')
card(s,8.88,1.95,3.86,3.1,'Desplegament','Publicació i configuració de l’aplicació sobre la infraestructura.')
text(s,'Sistemes informàtics aporta la base d’entorn i xarxa.\nEl projecte intermodular permet combinar aquests aprenentatges.',.65,5.45,12,1.1,23)

s = slide('Com treballarem: programarem junts', '03 · Desenvolupament del mòdul',
          'Metodologia confirmada pel professor el 2026-09-21; wiki/orientacio-docent.md.')
for x,n,t,b in [(.6,'01','Preparació','El professor prepara el projecte Spring amb les classes resoltes.'),(4.74,'02','Code along','Construïm el codi a classe i ens aturam per explicar els conceptes.'),(8.88,'03','Pràctica','Aplicau i repetiu el que hem treballat per preparar l’examen.')]:
    card(s,x,2,3.86,3.8,t,b,n)
banner(s,'Aquest mètode es manté a totes les unitats.')

s = slide('Quatre unitats, un recorregut', '03 · Desenvolupament del mòdul', unit_note(1))
for i,(t,b) in enumerate([('Introducció al servidor','Comprendre l’entorn i Spring.'),('Accés a dades','Connectar l’aplicació amb les dades.'),('Serveis web','Exposar i consumir funcionalitats.'),('Aplicacions web','Integrar les peces en una aplicació.')]):
    x=.6+(i%2)*6.23; y=1.95+(i//2)*2.25
    card(s,x,y,5.9,2.02,f'U{i+1} · {t}',b)

s = slide('U1 · Introducció a la programació\nen entorn servidor', '03 · Unitats i resultats d’aprenentatge', unit_note(1))
text(s,'Entendrem l’entorn servidor i com Spring crea i connecta els objectes: IoC, injecció de dependències i beans.',.6,2.03,12,1.25,25)
text(s,'RESULTAT D’APRENENTATGE',.6,3.7,9,.3,12,GREEN,True)
text(s,'PES GLOBAL',10.7,3.7,2,.3,12,GREEN,True,PP_ALIGN.RIGHT)
ra(s,'RA1',4.1,1.48,23)
banner(s,f'Durada · {H[0]} hores')

s = slide('U2 · Accés a dades', '03 · Unitats i resultats d’aprenentatge', unit_note(2))
text(s,'Connectarem l’aplicació amb les dades: consultar, crear, modificar i eliminar informació amb repositoris i serveis.',.6,1.95,12,1.25,25)
text(s,'RESULTAT D’APRENENTATGE',.6,3.7,9,.3,12,GREEN,True)
text(s,'PES GLOBAL',10.7,3.7,2,.3,12,GREEN,True,PP_ALIGN.RIGHT)
ra(s,'RA6',4.1,1.48,23)
banner(s,f'Durada · {H[1]} hores')

s = slide('U3 · Serveis web', '03 · Unitats i resultats d’aprenentatge', unit_note(3))
text(s,'Crearem i consumirem serveis web i reutilitzarem informació i components d’altres sistemes.',.6,1.9,12,.95,25)
text(s,'RESULTATS D’APRENENTATGE',.6,3.1,9,.3,12,GREEN,True)
text(s,'PES GLOBAL',10.7,3.1,2,.3,12,GREEN,True,PP_ALIGN.RIGHT)
ra(s,'RA7',3.5,1.13,22); ra(s,'RA9',4.78,1.13,22)
banner(s,f'Durada · {H[2]} hores')

s = slide('U4 · Aplicacions web', '03 · Unitats i resultats d’aprenentatge · 1/2', unit_note(4))
text(s,'Integrarem pàgines dinàmiques, formularis, estat, autenticació i interacció client-servidor amb una arquitectura per capes.',.6,1.9,12,1.15,25)
text(s,'RESULTATS D’APRENENTATGE',.6,3.25,9,.3,12,GREEN,True)
text(s,'PES GLOBAL',10.7,3.25,2,.3,12,GREEN,True,PP_ALIGN.RIGHT)
ra(s,'RA2',3.65,1.27,22); ra(s,'RA3',5.05,1.10,22)
text(s,'Continua amb RA4, RA5 i RA8 →',.6,6.34,12,.45,21,GREEN,True)

s = slide('U4 · Aplicacions web: continuació', '03 · Unitats i resultats d’aprenentatge · 2/2', unit_note(4))
text(s,'RESULTATS D’APRENENTATGE',.6,1.86,9,.3,12,GREEN,True)
text(s,'PES GLOBAL',10.7,1.86,2,.3,12,GREEN,True,PP_ALIGN.RIGHT)
ra(s,'RA4',2.20,1.30,22); ra(s,'RA5',3.65,1.06,22); ra(s,'RA8',4.86,1.06,22)
banner(s,f'Durada de tota U4 · {H[3]} hores')

evaluation = ('Font: full 4-AVALUACIÓ!A4:A10 i 2-ANCORATGE CURRICULAR!A4:F12. '
              'Precisió directa del professor: exàmens en paper, pràctiques per preparar-los i recuperació dels RA suspesos a final de curs. '
              'El llindar de 5 o més es pren del full de càlcul.')
s = slide('Avaluació: practicar i demostrar què sabem', '04 · Avaluació i qualificació', evaluation)
card(s,.6,2.0,5.9,3.62,'Pràctiques','Serveixen per aprendre i preparar l’examen.\n\nTenen ponderació 0% en la nota.')
card(s,6.83,2.0,5.9,3.62,'Exàmens en paper','Les proves individuals permeten valorar els resultats d’aprenentatge treballats.')
banner(s,'Les pràctiques obligatòries s’han de lliurar i aprovar.')

s = slide('Cal aprovar tots els resultats d’aprenentatge', '04 · Avaluació i qualificació', evaluation)
text(s,'Cada RA ≥ 5',.6,2.05,12,1.12,48,GREEN,True)
text(s,'Una bona nota en un RA no compensa un altre RA suspès.',.65,3.5,11.8,.88,27)
text(s,'Una nota de RA inferior a 5 no s’arrodoneix per convertir-la en un aprovat.',.65,4.63,11.8,1.05,25)
banner(s,'Cada prova individual ha de tenir almenys un 5 per fer mitjana.')

s = slide('L’arrodoniment és el darrer pas', '04 · Avaluació i qualificació', evaluation)
for x,n,t,b in [(.6,'01','Comprovar els RA','Tots han d’estar aprovats amb una nota igual o superior a 5.'),(4.74,'02','Aplicar els pesos','Calcular la nota ponderada amb els percentatges globals dels RA.'),(8.88,'03','Arrodonir la nota','Aplicar l’arrodoniment natural a la nota final del mòdul.')]:
    card(s,x,1.95,3.86,3.75,t,b,n)
banner(s,'L’arrodoniment només s’aplica després de superar tots els RA.')

s = slide('Tres exemples per entendre-ho', '04 · Avaluació i qualificació', evaluation+' Exemples numèrics il·lustratius.')
for x,t,b in [(.6,'Un RA té 4,9','El RA continua suspès.\n\nEncara que la mitjana sigui alta, cal recuperar-lo.'),(4.74,'Tots aprovats: 6,49','La nota ponderada final és 6,49.\n\nArrodoniment natural: 6.'),(8.88,'Tots aprovats: 6,50','La nota ponderada final és 6,50.\n\nArrodoniment natural: 7.')]:
    card(s,x,2.03,3.86,4.25,t,b)

s = slide('Recuperació dels RA suspesos', '04 · Avaluació i qualificació', evaluation)
text(s,'A final de curs',.6,2.0,12,.9,40,GREEN,True)
text(s,'Es podran recuperar els resultats d’aprenentatge que hagin quedat suspesos.',.65,3.1,11.8,.97,27)
text(s,'La recuperació comprovarà els aprenentatges no assolits.\nSegons la programació, la nota màxima de les proves de recuperació és 5.',.65,4.35,11.8,1.42,24)
banner(s,'Per superar el mòdul, tots els RA han de quedar aprovats.')

s = slide('Altres criteris que hem de tenir clars', '04 · Avaluació i qualificació', evaluation)
for y,t,b in [(1.92,'Absències a proves','Cal justificar-les dins el termini i pels motius admesos pel centre; en cas contrari, la prova es qualifica amb 0.'),(3.48,'Còpia o plagi','La qualificació de l’activitat o prova afectada és 0.'),(4.86,'Ús d’IA generativa','Només quan el professor l’autoritzi expressament per a una activitat i amb un ús traçable.')]:
    text(s,t,.6,y,3.05,.85,25,GREEN,True)
    text(s,b,3.98,y,8.70,1.22,23)

OUT.mkdir(parents=True, exist_ok=True)
dest = OUT / 'presentacio-modul-2026-2027.pptx'
prs.save(dest)
# Derivat llegible per revisar continguts i fonts sense un visor de diapositives.
lines = ['# Presentació del mòdul 2026–2027', '', 'Professor: David Pons', '']
for i, (s,t,n) in enumerate(transcript,1):
    lines += [f'## {i}. {t.replace(chr(10), " ")}', '']
    for sh in s.shapes:
        if sh.has_text_frame and sh.text.strip(): lines += [sh.text, '']
    lines += ['Font / notes del professor: '+n, '']
(OUT / 'contingut-i-fonts.md').write_text('\n'.join(lines),encoding='utf-8')
assert int(Decimal('6.50').quantize(Decimal('1'),rounding=ROUND_HALF_UP)) == 7
print(f'{dest}\n{len(prs.slides)} diapositives; {sum(H)} h; pesos RA: 100%.')
